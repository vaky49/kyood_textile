import csv
from decimal import Decimal, InvalidOperation
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict, Field, ValidationError

from app.domain.requirements import ProductCategory


class CatalogueRow(BaseModel):
    model_config = ConfigDict(extra="ignore", str_strip_whitespace=True)

    title: str = Field(min_length=1, max_length=300)
    category: ProductCategory
    description: str | None = None
    sku: str | None = Field(default=None, max_length=100)
    subcategory: str | None = Field(default=None, max_length=100)
    fabric: str | None = Field(default=None, max_length=100)
    color: str | None = Field(default=None, max_length=100)
    style: str | None = Field(default=None, max_length=200)
    price: Decimal | None = Field(default=None, ge=0)
    currency: str | None = Field(default=None, min_length=3, max_length=3)
    moq: int | None = Field(default=None, gt=0)
    stock_status: str | None = Field(default=None, max_length=50)
    available_quantity: int | None = Field(default=None, ge=0)


class CatalogueRowError(BaseModel):
    row_number: int
    detail: str


class CatalogueParseResult(BaseModel):
    rows: list[CatalogueRow]
    errors: list[CatalogueRowError]
    total_rows: int


def parse_catalogue_spreadsheet(filename: str, content: bytes) -> CatalogueParseResult:
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        table = _read_csv(content)
    elif suffix in {".xlsx", ".xlsm"}:
        table = _read_excel(content)
    else:
        raise ValueError("Only CSV and XLSX spreadsheets can be parsed automatically")

    if not table:
        return CatalogueParseResult(rows=[], errors=[], total_rows=0)
    headers = [_normalize_header(value) for value in table[0]]
    rows: list[CatalogueRow] = []
    errors: list[CatalogueRowError] = []
    for row_number, values in enumerate(table[1:], start=2):
        if not any(value not in (None, "") for value in values):
            continue
        try:
            record = {
                header: value
                for header, value in zip(headers, values, strict=False)
                if header is not None
            }
            record = _normalize_record(record)
            if isinstance(record.get("category"), str):
                record["category"] = _normalize_category(record["category"])
            rows.append(CatalogueRow.model_validate(record))
        except (InvalidOperation, ValidationError, ValueError) as error:
            detail = str(error).splitlines()[-1][:400]
            errors.append(CatalogueRowError(row_number=row_number, detail=detail))
    return CatalogueParseResult(rows=rows, errors=errors, total_rows=len(table) - 1)


def _read_csv(content: bytes) -> list[list[Any]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise ValueError("CSV must use UTF-8 encoding") from error
    return list(csv.reader(StringIO(text)))


def _read_excel(content: bytes) -> list[list[Any]]:
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as error:
        raise ValueError("XLSX file could not be opened") from error
    worksheet = workbook.active
    values = [list(row) for row in worksheet.iter_rows(values_only=True)]
    workbook.close()
    return values


def _normalize_header(value: Any) -> str | None:
    if value is None:
        return None
    normalized = "_".join(str(value).strip().casefold().replace("/", " ").split())
    aliases = {
        "product": "title",
        "item": "title",
        "name": "title",
        "product_name": "title",
        "item_name": "title",
        "sub_category": "subcategory",
        "minimum_order_quantity": "moq",
        "minimum_quantity": "moq",
        "qty_available": "available_quantity",
        "stock_quantity": "available_quantity",
        "rate": "price",
        "price_per_piece": "price",
        "fabric_type": "fabric",
    }
    return aliases.get(normalized, normalized)


def _normalize_record(record: dict[str, Any]) -> dict[str, Any]:
    for field in ("title", "category", "description", "sku", "subcategory", "fabric", "color", "style", "currency", "stock_status"):
        value = record.get(field)
        if isinstance(value, str) and not value.strip():
            record[field] = None
    currency = record.get("currency")
    if isinstance(currency, str):
        record["currency"] = currency.upper()
    price = record.get("price")
    if isinstance(price, str):
        cleaned = price.replace(",", "").replace("₹", "").replace("Rs.", "").strip()
        if cleaned:
            record["price"] = Decimal(cleaned)
        else:
            record["price"] = None
    for field in ("moq", "available_quantity"):
        value = record.get(field)
        if isinstance(value, str):
            cleaned = value.replace(",", "").strip()
            record[field] = int(cleaned) if cleaned else None
    return record


def _normalize_category(value: str) -> ProductCategory:
    normalized = value.strip().casefold().replace(" ", "_").replace("-", "_")
    aliases = {
        "suit": ProductCategory.LADIES_SUIT,
        "ladies_suits": ProductCategory.LADIES_SUIT,
        "salwar_suits": ProductCategory.SALWAR_SUIT,
        "dress_materials": ProductCategory.DRESS_MATERIAL,
    }
    if normalized in aliases:
        return aliases[normalized]
    return ProductCategory(normalized)